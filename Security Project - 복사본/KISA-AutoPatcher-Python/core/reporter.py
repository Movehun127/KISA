# -*- coding: utf-8 -*-
"""
reporter.py  –  보고서 생성 엔진 (PowerShell Reporter.psm1 대체)
Markdown + Excel(openpyxl) 리포트를 생성합니다.
"""

import os
import sys
import shutil
from datetime import datetime
import openpyxl


# ── 엑셀 템플릿 경로 설정 (상대 경로 변경) ───────────────────────────
if getattr(sys, 'frozen', False):
    _base_dir = os.path.dirname(sys.executable)
else:
    _base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

from .paths import RESOURCE_DIR
EXCEL_TEMPLATE = os.path.join(RESOURCE_DIR, "config", "Model2.xlsx")


def invoke_reporting(
    initial_results: list[dict],
    final_results: list[dict],
    report_dir: str,
    evidence_dir: str = "",
    log_callback=None
) -> None:
    os.makedirs(report_dir, exist_ok=True)
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    pc_name   = os.environ.get("COMPUTERNAME", "UNKNOWN")

    def log(msg, level="info"):
        if log_callback:
            log_callback(msg, level)

    # ── 1. Markdown 리포트 생성 ───────────────────────────────────────
    md_path = os.path.join(report_dir, "Security_Patch_Report.md")
    lines = [
        f"\n---",
        f"## 🛡️ KISA 보안 패치 보고서",
        f"**생성 일시:** {timestamp}",
        f"**대상 PC:** {pc_name}",
        "",
        "| 항목 코드 | 항목명 | 중요도 | 조치 전 상태 | 조치 후 상태 | 조치 결과 | 증빙 |",
        "|---|---|---|---|---|---|---|",
    ]

    final_map = {r["ItemId"]: r for r in final_results}

    for init in initial_results:
        iid    = init["ItemId"]
        final  = final_map.get(iid, init)
        is_ok  = final["Status"] == "양호"
        icon   = "✅" if is_ok else ("⚠️" if "수동" in final["Status"] else "❌")

        init_str = f"{init['Status']} ({init.get('CurrentValue', '')})"
        final_str = f"{final['Status']} ({final.get('CurrentValue', '')})"

        cells = [iid, init['Title'], init['Level'], init_str, final_str,
                 init.get('FixStatus', '미실행'), init.get('EvidenceStatus', '미수집')]
        lines.append('| ' + ' | '.join(str(c).replace('|', '\\|').replace('\r', '').replace('\n', '<br>') for c in cells) + ' |')

    lines.append("")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    log(f"✅ 마크다운 보고서 생성: {md_path}", "good")

    # ── 2. Excel 리포트 생성 ─────────────────────────────────────────
    if not os.path.exists(EXCEL_TEMPLATE):
        log(f"⚠️ 엑셀 템플릿 없음: {EXCEL_TEMPLATE}", "warn")
        return

    excel_path = os.path.join(report_dir, "Security_Patch_Report.xlsx")
    shutil.copy2(EXCEL_TEMPLATE, excel_path)

    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active

        for row in ws.iter_rows(min_row=6):
            code_cell = row[3]  # 4번째 열 = 항목코드
            code = str(code_cell.value or "").strip()
            if not code.startswith("W-"):
                continue

            # 해당 코드에 매칭되는 final_results 수집
            matched = [r for r in final_results
                       if r["ItemId"] == code or r["ItemId"].startswith(code + "_")]
            if not matched:
                continue

            # 미사용 서비스 체크
            unused_items = [m for m in matched if "미사용" in m.get("Status", "")]
            if unused_items:
                row[4].value = "X"   # 조치 여부 열 (5번째)
                row[5].value = unused_items[0]["Status"].replace("양호(", "").rstrip(")")
                continue

            if all(m.get("TechType") == "Type_Skip" for m in matched):
                row[4].value = "확인 필요"
                row[5].value = "수동 검토 항목: 자동 적합 판정 없음"
                continue

            has_vuln   = any("양호" not in m["Status"] and "수동 조치" not in m["Status"]
                             for m in matched)
            has_manual = any("수동 조치" in m["Status"] for m in matched)
            val_strs   = []

            for m in matched:
                if "수동 조치" in m["Status"]:
                    val_strs.append(f"{m['ItemId']}: 수동 조치 필요")
                elif "양호" not in m["Status"]:
                    val_strs.append(f"{m['ItemId']}: 취약(현재값={m.get('CurrentValue', '')})")

            row[4].value = "X" if has_vuln else ("확인 필요" if has_manual else "O")   # 조치 여부
            row[5].value = ", ".join(val_strs) if val_strs else ""  # 비고

        if '실행 증빙' in wb.sheetnames:
            del wb['실행 증빙']
        evidence_sheet = wb.create_sheet('실행 증빙')
        evidence_sheet.append(['항목', '조치 결과', '재점검 결과', '화면 캡처', '증빙 JSON', '비고'])
        for init in initial_results:
            evidence_sheet.append([init['ItemId'], init.get('FixStatus', '미실행'),
                final_map.get(init['ItemId'], {}).get('Status', '확인 불가'),
                init.get('EvidenceStatus', '미수집'), init.get('EvidenceFile', ''),
                init.get('Note', '')])
        wb.save(excel_path)
        log(f"✅ 엑셀 보고서 생성: {excel_path}", "good")

    except Exception as e:
        log(f"⚠️ 엑셀 보고서 생성 오류: {e}", "warn")
